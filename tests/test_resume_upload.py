"""
Resume upload (PDF/DOCX -> text extraction -> existing ResumeParserAgent ->
auto-fill preview).

Covers `utils/resume_documents.py` (pure text extraction, no LLM/agent
involved) and the new `POST /candidates/parse-resume-file` endpoint
(api/routes/candidates.py), which reuses the EXISTING mock-provider
`ResumeParserAgent` exactly the way `POST /candidates` already does (see
tests/test_backend_jobs_candidates_applications.py's own module docstring
for that convention) - no second parser, no fabricated result.

`tests/fixtures/sample_resume.pdf` / `.docx` are REAL files with real
embedded resume text (generated once via `fpdf2`/`python-docx`, not
project runtime dependencies - see their own content), satisfying "use
real files, not hardcoded fake parsed results" for the extraction layer.
The LLM step itself runs against MockLLMProvider like the rest of this
suite; extraction correctness is what these files exist to prove.
"""
import os

import pytest
from fastapi.testclient import TestClient

from api.app import create_app
from core.errors import DependencyError
from repositories.memory import InMemoryCandidateRepository
from services.candidate_service import CandidateService
from utils.resume_documents import (
    ResumeExtractionError,
    UnsupportedResumeFormatError,
    extract_resume_text,
)

_FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures")
_PDF_PATH = os.path.join(_FIXTURES_DIR, "sample_resume.pdf")
_DOCX_PATH = os.path.join(_FIXTURES_DIR, "sample_resume.docx")
# Genuinely image-only documents (a rendered resume-text PNG embedded with
# no real text layer/paragraph text at all - see how these were generated
# in the OCR implementation task) - used to prove the OCR fallback path,
# not the normal-extraction path.
_SCANNED_PDF_PATH = os.path.join(_FIXTURES_DIR, "scanned_resume.pdf")
_SCANNED_DOCX_PATH = os.path.join(_FIXTURES_DIR, "scanned_resume.docx")


def _read(path: str) -> bytes:
    with open(path, "rb") as f:
        return f.read()


@pytest.fixture
def client():
    with TestClient(create_app()) as c:
        yield c


# ---------------------------------------------------------------------------
# utils/resume_documents.py - pure extraction, no agent/LLM involved
# ---------------------------------------------------------------------------

class TestExtractResumeText:
    def test_real_pdf_extracts_real_text(self):
        text, source_format = extract_resume_text("sample_resume.pdf", _read(_PDF_PATH))
        assert source_format == "pdf"
        assert "Morgan Reyes" in text
        assert "PostgreSQL" in text

    def test_real_docx_extracts_real_text(self):
        text, source_format = extract_resume_text("sample_resume.docx", _read(_DOCX_PATH))
        assert source_format == "docx"
        assert "Priya Nair" in text
        assert "React" in text

    def test_unsupported_extension_is_rejected(self):
        with pytest.raises(UnsupportedResumeFormatError):
            extract_resume_text("resume.txt", b"plain text resume")

    def test_legacy_doc_is_explicitly_rejected_not_silently_ignored(self):
        with pytest.raises(UnsupportedResumeFormatError, match=r"\.docx or \.pdf"):
            extract_resume_text("resume.doc", b"not a real doc file")

    def test_corrupted_pdf_raises_extraction_error(self):
        with pytest.raises(ResumeExtractionError):
            extract_resume_text("resume.pdf", b"this is not a real pdf file")

    def test_corrupted_docx_raises_extraction_error(self):
        with pytest.raises(ResumeExtractionError):
            extract_resume_text("resume.docx", b"this is not a real docx file")

    def test_scanned_pdf_falls_back_to_ocr_and_extracts_real_text(self):
        """A genuinely image-only PDF (no text layer at all - confirmed by
        pypdf extracting nothing from it) must still produce the resume's
        actual content, via OCR - not an empty/fabricated result."""
        text, source_format = extract_resume_text(
            "scanned_resume.pdf", _read(_SCANNED_PDF_PATH)
        )
        assert source_format == "pdf"
        assert "Jamie Okafor" in text or "Okafor" in text
        assert "Data Engineer" in text
        assert "PostgreSQL" in text or "Postgres" in text  # OCR may miss exact casing/spacing

    def test_image_based_docx_falls_back_to_ocr_and_extracts_real_text(self):
        """A DOCX containing only an embedded resume-text image, no real
        paragraph text - python-docx alone extracts nothing from it."""
        text, source_format = extract_resume_text(
            "scanned_resume.docx", _read(_SCANNED_DOCX_PATH)
        )
        assert source_format == "docx"
        assert "Okafor" in text
        assert "Apache Spark" in text

    def test_normal_pdf_does_not_invoke_ocr_at_all(self, monkeypatch):
        """OCR must not be forced on every document - a normal, text-based
        PDF should never even construct the OCR engine."""
        import utils.resume_documents as module

        def _fail_if_called(*args, **kwargs):
            raise AssertionError("OCR must not run when normal extraction already succeeded")

        monkeypatch.setattr(module, "_get_ocr_engine", _fail_if_called)
        text, _ = extract_resume_text("sample_resume.pdf", _read(_PDF_PATH))
        assert "Morgan Reyes" in text

    def test_normal_docx_does_not_invoke_ocr_at_all(self, monkeypatch):
        import utils.resume_documents as module

        def _fail_if_called(*args, **kwargs):
            raise AssertionError("OCR must not run when normal extraction already succeeded")

        monkeypatch.setattr(module, "_get_ocr_engine", _fail_if_called)
        text, _ = extract_resume_text("sample_resume.docx", _read(_DOCX_PATH))
        assert "Priya Nair" in text

    def test_docx_with_no_text_raises_extraction_error(self):
        from docx import Document
        import io

        doc = Document()
        doc.add_paragraph("")
        doc.add_paragraph("   ")
        buf = io.BytesIO()
        doc.save(buf)
        with pytest.raises(ResumeExtractionError, match="No extractable text"):
            extract_resume_text("blank.docx", buf.getvalue())


# ---------------------------------------------------------------------------
# CandidateService.preview_resume - parses via the EXISTING ResumeParserAgent
# without persisting (services/candidate_service.py)
# ---------------------------------------------------------------------------

class TestPreviewResume:
    @pytest.mark.asyncio
    async def test_preview_does_not_persist_a_candidate(self):
        repo = InMemoryCandidateRepository()
        service = CandidateService(candidate_repository=repo)

        parsed_resume, used_fallback, warning = await service.preview_resume(
            resume_text="Jane Doe\nSkills: Python, SQL.", candidate_name="Jane Doe",
        )

        assert parsed_resume.candidate_name == "Jane Doe"
        assert await repo.list_candidates() == []

    @pytest.mark.asyncio
    async def test_preview_parser_failure_is_a_dependency_error(self):
        class _BrokenParser:
            async def run(self, **kwargs):
                return {"result": {"parsed_resume": None, "error": "even the fallback failed"}}

        service = CandidateService(
            candidate_repository=InMemoryCandidateRepository(),
            resume_parser_factory=lambda: _BrokenParser(),
        )
        with pytest.raises(DependencyError):
            await service.preview_resume(resume_text="Jane Doe\nSkills: Python.", candidate_name="Jane Doe")


# ---------------------------------------------------------------------------
# POST /candidates/parse-resume-file
# ---------------------------------------------------------------------------

class TestParseResumeFileEndpoint:
    def test_real_pdf_upload_returns_structured_preview(self, client):
        with open(_PDF_PATH, "rb") as f:
            response = client.post(
                "/candidates/parse-resume-file",
                files={"resume_file": ("sample_resume.pdf", f, "application/pdf")},
                data={"candidate_name": "Morgan Reyes"},
            )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["parsed_resume"]["candidate_name"] == "Morgan Reyes"
        assert body["parsed_resume"]["source_format"] == "pdf"
        assert "PostgreSQL" in body["resume_text"]
        assert isinstance(body["used_fallback"], bool)

    def test_real_docx_upload_returns_structured_preview(self, client):
        with open(_DOCX_PATH, "rb") as f:
            response = client.post(
                "/candidates/parse-resume-file",
                files={
                    "resume_file": (
                        "sample_resume.docx", f,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"candidate_name": "Priya Nair"},
            )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["parsed_resume"]["candidate_name"] == "Priya Nair"
        assert body["parsed_resume"]["source_format"] == "docx"
        assert "React" in body["resume_text"]

    def test_scanned_pdf_upload_ocrs_and_the_real_text_reaches_the_agent(self, client):
        """End-to-end proof of the OCR fallback THROUGH the actual endpoint:
        a genuinely image-only PDF still produces a real structured preview,
        and `resume_text` (what ResumeParserAgent actually received) is the
        OCR'd content, not empty/fabricated."""
        with open(_SCANNED_PDF_PATH, "rb") as f:
            response = client.post(
                "/candidates/parse-resume-file",
                files={"resume_file": ("scanned_resume.pdf", f, "application/pdf")},
                data={"candidate_name": "Jamie Okafor"},
            )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["parsed_resume"]["source_format"] == "pdf"
        assert "Okafor" in body["resume_text"]
        assert "Apache Spark" in body["resume_text"]

    def test_scanned_docx_upload_ocrs_and_the_real_text_reaches_the_agent(self, client):
        with open(_SCANNED_DOCX_PATH, "rb") as f:
            response = client.post(
                "/candidates/parse-resume-file",
                files={
                    "resume_file": (
                        "scanned_resume.docx", f,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"candidate_name": "Jamie Okafor"},
            )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["parsed_resume"]["source_format"] == "docx"
        assert "Okafor" in body["resume_text"]
        assert "PostgreSQL" in body["resume_text"] or "Postgres" in body["resume_text"]

    def test_unsupported_file_type_is_a_clean_400(self, client):
        response = client.post(
            "/candidates/parse-resume-file",
            files={"resume_file": ("resume.txt", b"plain text", "text/plain")},
            data={"candidate_name": "Test User"},
        )
        assert response.status_code == 400
        body = response.json()
        assert body["error"] == "invalid_request"
        assert "traceback" not in body["detail"].lower()

    def test_legacy_doc_file_is_a_clean_400_naming_docx_as_the_alternative(self, client):
        response = client.post(
            "/candidates/parse-resume-file",
            files={"resume_file": ("resume.doc", b"old format", "application/msword")},
            data={"candidate_name": "Test User"},
        )
        assert response.status_code == 400
        assert ".docx" in response.json()["detail"]

    def test_corrupted_pdf_is_a_clean_400_not_a_traceback(self, client):
        response = client.post(
            "/candidates/parse-resume-file",
            files={"resume_file": ("resume.pdf", b"not a real pdf", "application/pdf")},
            data={"candidate_name": "Test User"},
        )
        assert response.status_code == 400
        assert response.json()["error"] == "invalid_request"

    def test_empty_file_is_a_clean_400(self, client):
        response = client.post(
            "/candidates/parse-resume-file",
            files={"resume_file": ("resume.pdf", b"", "application/pdf")},
            data={"candidate_name": "Test User"},
        )
        assert response.status_code == 400

    def test_missing_candidate_name_is_a_422(self, client):
        with open(_PDF_PATH, "rb") as f:
            response = client.post(
                "/candidates/parse-resume-file",
                files={"resume_file": ("sample_resume.pdf", f, "application/pdf")},
            )
        assert response.status_code == 422

    def test_upload_does_not_create_a_candidate_record(self, client):
        """The whole point of this endpoint (A5) is that the candidate
        reviews/edits before anything is persisted - confirm the existing
        candidate list stays empty after a parse-only call."""
        with open(_PDF_PATH, "rb") as f:
            client.post(
                "/candidates/parse-resume-file",
                files={"resume_file": ("sample_resume.pdf", f, "application/pdf")},
                data={"candidate_name": "Morgan Reyes"},
            )
        listing = client.get("/candidates").json()
        assert listing["total"] == 0

    def test_confirmed_preview_can_then_be_registered_via_the_existing_endpoint(self, client):
        """The end of Feature A's flow: once the candidate confirms the
        (possibly edited) extracted text, the EXISTING POST /candidates is
        what actually creates the record - unmodified, no second API."""
        with open(_PDF_PATH, "rb") as f:
            preview = client.post(
                "/candidates/parse-resume-file",
                files={"resume_file": ("sample_resume.pdf", f, "application/pdf")},
                data={"candidate_name": "Morgan Reyes"},
            ).json()

        created = client.post(
            "/candidates",
            json={"candidate_name": "Morgan Reyes", "resume_text": preview["resume_text"]},
        )
        assert created.status_code == 201, created.text
        candidate_id = created.json()["candidate_id"]

        fetched = client.get(f"/candidates/{candidate_id}")
        assert fetched.status_code == 200
        assert fetched.json()["resume"]["candidate_name"] == "Morgan Reyes"
